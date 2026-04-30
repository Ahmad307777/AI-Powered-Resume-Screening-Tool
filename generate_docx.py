import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def convert_md_to_docx(md_path, docx_path):
    doc = Document()
    
    # Set default style to 'Normal' and font to Arial
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    for line in lines:
        line = line.strip()
        if not line:
            doc.add_paragraph()
            continue
            
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2)
        elif line.startswith('* ') or line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        else:
            p = doc.add_paragraph()
            # Basic bolding handling
            parts = line.split('**')
            for i, part in enumerate(parts):
                run = p.add_run(part)
                if i % 2 == 1:
                    run.bold = True

    doc.save(docx_path)
    print(f"File saved to {docx_path}")

if __name__ == "__main__":
    md_file = r'c:\Users\namal\Downloads\Resume-Screening-main\Full_Project_Analysis_and_Technical_Report.md'
    docx_file = r'c:\Users\namal\Downloads\Resume-Screening-main\Full_Project_Analysis_and_Technical_Report.docx'
    convert_md_to_docx(md_file, docx_file)
