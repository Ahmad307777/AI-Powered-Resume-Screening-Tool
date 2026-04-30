import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def convert_md_to_docx(md_path, docx_path):
    doc = Document()
    
    # Set default style to 'Normal' and font to Times New Roman if possible
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        if line.startswith('# '):
            p = doc.add_heading(line[2:], level=0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            doc.add_heading(line[3:], level=1)
        elif line.startswith('### '):
            doc.add_heading(line[4:], level=2)
        elif line.startswith('---'):
            # Just ignore horizontal rules or add a page break? 
            # In a research paper, maybe just skip or add space.
            pass
        elif line.startswith('* ') or line.startswith('- '):
            doc.add_paragraph(line[2:], style='List Bullet')
        else:
            # Handle basic bolding like **text**
            p = doc.add_paragraph()
            parts = line.split('**')
            for i, part in enumerate(parts):
                run = p.add_run(part)
                if i % 2 == 1:
                    run.bold = True

    doc.save(docx_path)
    print(f"File saved to {docx_path}")

if __name__ == "__main__":
    md_file = r'c:\Users\namal\Downloads\Resume-Screening-main\Milestone_1_Proposal.md'
    docx_file = r'c:\Users\namal\Downloads\Resume-Screening-main\Milestone_1_Research_Proposal.docx'
    convert_md_to_docx(md_file, docx_file)
